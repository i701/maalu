from datetime import date
from decimal import Decimal

from app.enums import ACCENT_COLORS
from httpx import AsyncClient

from tests.conftest import Signup

CAT = "/api/v1/categories"
TX = "/api/v1/transactions"
BUDGET = "/api/v1/budget"
DAY = date.today().isoformat()


async def make(client, auth, name, **extra) -> dict:
    resp = await client.post(CAT, json={"name": name, **extra}, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def spend(client, auth, category, amount, **extra) -> dict:
    body = {
        "type": "expense",
        "amount": amount,
        "payment_method": "cash",
        "category": category,
        "transaction_date": DAY,
        **extra,
    }
    resp = await client.post(TX, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def all_categories(client, auth) -> list[dict]:
    return (await client.get(CAT, headers=auth)).json()


# --- colours ------------------------------------------------


async def test_a_new_category_gets_the_colour_used_least(client: AsyncClient, auth):
    colours = [(await make(client, auth, f"Cat {n:02d}"))["color"] for n in range(10)]
    assert colours[:8] == list(ACCENT_COLORS)  # each colour once before any repeats
    assert colours[8:] == list(ACCENT_COLORS[:2])


async def test_categories_made_by_a_transaction_get_a_colour_too(client: AsyncClient, auth):
    t = await spend(client, auth, "Coffee", "5")
    assert t["category_color"] in ACCENT_COLORS
    [coffee] = await all_categories(client, auth)
    assert coffee["color"] == t["category_color"]


async def test_a_colour_can_be_chosen_and_changed(client: AsyncClient, auth):
    made = await make(client, auth, "Rent", color="pink")
    assert made["color"] == "pink"
    changed = await client.patch(f"{CAT}/{made['id']}", json={"color": "lime"}, headers=auth)
    assert changed.json()["color"] == "lime"
    assert changed.json()["name"] == "Rent"  # nothing else moved


async def test_only_the_palette_colours_are_accepted(client: AsyncClient, auth):
    assert (
        await client.post(CAT, json={"name": "X", "color": "red"}, headers=auth)
    ).status_code == 422
    made = await make(client, auth, "Y")
    for body in ({"color": "teal"}, {"color": None}, {"name": None}):
        resp = await client.patch(f"{CAT}/{made['id']}", json=body, headers=auth)
        assert resp.status_code == 422, body


async def test_colours_show_on_transactions_items_and_totals(client: AsyncClient, auth):
    await make(client, auth, "Groceries", color="green")
    await make(client, auth, "Gadgets", color="violet")
    t = await spend(
        client,
        auth,
        None,
        "100",
        splits=[
            {"name": "Eggs", "category": "Groceries", "amount": "60"},
            {"name": "Phones", "category": "Gadgets", "amount": "40"},
        ],
    )
    assert [s["category_color"] for s in t["splits"]] == ["green", "violet"]
    totals = (await client.get(f"{TX}/categories", headers=auth)).json()["categories"]
    assert {c["category"]: c["color"] for c in totals} == {
        "Groceries": "green",
        "Gadgets": "violet",
    }


async def test_colours_are_counted_per_user(client: AsyncClient, auth, signup: Signup):
    bob = await signup("bob@example.com")
    for n in range(3):
        await make(client, bob, f"B{n}")
    assert (await make(client, auth, "Mine"))["color"] == ACCENT_COLORS[0]


# --- parents ------------------------------------------------


async def test_a_sub_category_names_its_parent(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    groceries = await make(client, auth, "Groceries", parent_id=food["id"])
    assert groceries["parent_id"] == food["id"]
    assert groceries["parent"] == "Food"
    assert groceries["child_count"] == 0
    [reloaded] = [c for c in await all_categories(client, auth) if c["id"] == food["id"]]
    assert reloaded["child_count"] == 1
    assert reloaded["parent_id"] is None


async def test_the_parent_rules(client: AsyncClient, auth, signup: Signup):
    food = await make(client, auth, "Food")
    groceries = await make(client, auth, "Groceries", parent_id=food["id"])
    other = await make(client, auth, "Transport")

    # one level only: a sub-category can't be a parent
    nested = await client.post(
        CAT, json={"name": "Eggs", "parent_id": groceries["id"]}, headers=auth
    )
    assert nested.status_code == 422
    assert "sub-category itself" in nested.json()["detail"]
    # not itself
    own = await client.patch(f"{CAT}/{food['id']}", json={"parent_id": food["id"]}, headers=auth)
    assert own.status_code == 422
    # a parent with sub-categories can't go under another category
    top = await client.patch(f"{CAT}/{food['id']}", json={"parent_id": other["id"]}, headers=auth)
    assert top.status_code == 422
    assert "has sub-categories" in top.json()["detail"]
    # not someone else's, not a made-up one
    bob = await signup("bob@example.com")
    theirs = await make(client, bob, "Theirs")
    foreign = await client.patch(
        f"{CAT}/{other['id']}", json={"parent_id": theirs["id"]}, headers=auth
    )
    assert foreign.status_code == 422
    missing = await client.post(
        CAT, json={"name": "Z", "parent_id": "00000000-0000-0000-0000-000000000000"}, headers=auth
    )
    assert missing.status_code == 422


async def test_a_category_can_move_under_a_parent_and_back(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    cafe = await make(client, auth, "Cafe")
    moved = await client.patch(f"{CAT}/{cafe['id']}", json={"parent_id": food["id"]}, headers=auth)
    assert moved.json()["parent"] == "Food"
    back = await client.patch(f"{CAT}/{cafe['id']}", json={"parent_id": None}, headers=auth)
    assert back.json()["parent_id"] is None
    assert back.json()["parent"] is None
    # renaming leaves the parent alone
    renamed = await client.patch(
        f"{CAT}/{cafe['id']}", json={"parent_id": food["id"]}, headers=auth
    )
    renamed = await client.patch(f"{CAT}/{cafe['id']}", json={"name": "Cafes"}, headers=auth)
    assert renamed.json()["parent"] == "Food"


async def test_the_list_keeps_a_parent_with_its_children(client: AsyncClient, auth):
    transport = await make(client, auth, "Transport")
    food = await make(client, auth, "Food")
    await make(client, auth, "Taxi", parent_id=transport["id"])
    await make(client, auth, "Dining", parent_id=food["id"])
    await make(client, auth, "Groceries", parent_id=food["id"])
    await make(client, auth, "Bus", parent_id=transport["id"])
    await make(client, auth, "Health")
    names = [c["name"] for c in await all_categories(client, auth)]
    assert names == ["Food", "Dining", "Groceries", "Health", "Transport", "Bus", "Taxi"]


async def test_deleting_a_parent_leaves_its_children_top_level(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    groceries = await make(client, auth, "Groceries", parent_id=food["id"])
    assert (await client.delete(f"{CAT}/{food['id']}", headers=auth)).status_code == 204
    [left] = await all_categories(client, auth)
    assert left["id"] == groceries["id"]
    assert left["parent_id"] is None


async def test_a_parent_that_is_used_cannot_be_deleted(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    await make(client, auth, "Groceries", parent_id=food["id"])
    await spend(client, auth, "Food", "10")
    assert (await client.delete(f"{CAT}/{food['id']}", headers=auth)).status_code == 409


async def test_merging_a_sub_category_into_another_category(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    cafe = await make(client, auth, "Cafe", parent_id=food["id"])
    eats = await make(client, auth, "Eats", parent_id=food["id"])
    await spend(client, auth, "Cafe", "10")
    resp = await client.post(
        f"{CAT}/{cafe['id']}/merge", json={"into_id": eats["id"]}, headers=auth
    )
    assert resp.status_code == 200
    assert resp.json()["parent"] == "Food"
    assert resp.json()["transaction_count"] == 1


async def test_merging_a_parent_moves_its_children_under_the_target(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    await make(client, auth, "Cafe", parent_id=food["id"])
    eating = await make(client, auth, "Eating out")
    resp = await client.post(
        f"{CAT}/{food['id']}/merge", json={"into_id": eating["id"]}, headers=auth
    )
    assert resp.status_code == 200
    cafe = next(c for c in await all_categories(client, auth) if c["name"] == "Cafe")
    assert cafe["parent_id"] == eating["id"]


async def test_a_parent_cannot_be_merged_into_a_sub_category_of_another(client: AsyncClient, auth):
    food = await make(client, auth, "Food")
    await make(client, auth, "Cafe", parent_id=food["id"])
    other = await make(client, auth, "Transport")
    taxi = await make(client, auth, "Taxi", parent_id=other["id"])
    resp = await client.post(
        f"{CAT}/{food['id']}/merge", json={"into_id": taxi["id"]}, headers=auth
    )
    assert resp.status_code == 422
    assert "sub-category" in resp.json()["detail"]


async def test_merging_a_parent_into_its_own_child_lets_the_child_take_its_place(
    client: AsyncClient, auth
):
    food = await make(client, auth, "Food")
    cafe = await make(client, auth, "Cafe", parent_id=food["id"])
    resp = await client.post(
        f"{CAT}/{food['id']}/merge", json={"into_id": cafe["id"]}, headers=auth
    )
    assert resp.status_code == 200
    assert resp.json()["parent_id"] is None


# --- a parent includes its children ------------------------------------------------


async def family(client, auth):
    food = await make(client, auth, "Food")
    await make(client, auth, "Groceries", parent_id=food["id"])
    await make(client, auth, "Dining", parent_id=food["id"])
    await make(client, auth, "Fuel")
    return food


async def test_filtering_by_a_parent_includes_its_children(client: AsyncClient, auth):
    food = await family(client, auth)
    await spend(client, auth, "Food", "10")
    await spend(client, auth, "Groceries", "20")
    await spend(client, auth, "Dining", "30")
    await spend(client, auth, "Fuel", "40")
    rows = (await client.get(TX, params={"category_id": food["id"]}, headers=auth)).json()
    assert sorted(r["category"] for r in rows) == ["Dining", "Food", "Groceries"]
    totals = (
        await client.get(f"{TX}/summary", params={"category_id": food["id"]}, headers=auth)
    ).json()
    assert Decimal(totals["expense"]) == 60
    ids = (await client.get(f"{TX}/ids", params={"category_id": food["id"]}, headers=auth)).json()
    assert len(ids) == 3


async def test_filtering_by_a_child_is_just_the_child(client: AsyncClient, auth):
    await family(client, auth)
    groceries = next(c for c in await all_categories(client, auth) if c["name"] == "Groceries")
    await spend(client, auth, "Food", "10")
    await spend(client, auth, "Groceries", "20")
    rows = (await client.get(TX, params={"category_id": groceries["id"]}, headers=auth)).json()
    assert [r["category"] for r in rows] == ["Groceries"]


async def test_a_parent_filter_counts_the_children_items_of_split_purchases(
    client: AsyncClient, auth
):
    food = await family(client, auth)
    await spend(
        client,
        auth,
        None,
        "100",
        splits=[
            {"name": "Eggs", "category": "Groceries", "amount": "60"},
            {"name": "Petrol", "category": "Fuel", "amount": "40"},
        ],
    )
    totals = (
        await client.get(f"{TX}/summary", params={"category_id": food["id"]}, headers=auth)
    ).json()
    assert Decimal(totals["expense"]) == 60  # only the groceries item, which is under Food
    rows = (
        await client.get(f"{TX}/categories", params={"category_id": food["id"]}, headers=auth)
    ).json()["categories"]
    assert [(r["category"], r["parent"]) for r in rows] == [("Groceries", "Food")]


async def test_category_totals_say_which_parent_a_row_belongs_to(client: AsyncClient, auth):
    await family(client, auth)
    await spend(client, auth, "Groceries", "20")
    await spend(client, auth, "Fuel", "5")
    rows = (await client.get(f"{TX}/categories", headers=auth)).json()["categories"]
    by_name = {r["category"]: r for r in rows}
    assert by_name["Groceries"]["parent"] == "Food"
    assert by_name["Groceries"]["parent_id"] is not None
    assert by_name["Fuel"]["parent"] is None


async def test_the_export_filter_includes_a_parents_children(client: AsyncClient, auth):
    import io

    from openpyxl import load_workbook

    food = await family(client, auth)
    await spend(client, auth, "Groceries", "20")
    await spend(client, auth, "Fuel", "40")
    resp = await client.get(f"{TX}/export", params={"category_id": food["id"]}, headers=auth)
    wb = load_workbook(io.BytesIO(resp.content))
    rows = list(wb["Transactions"].iter_rows(values_only=True))
    assert [r[3] for r in rows[1:]] == ["Groceries"]


# --- budgets ------------------------------------------------


async def line(client, auth, name, category_ids, amount="500"):
    return await client.post(
        f"{BUDGET}/lines",
        json={"name": name, "category_ids": category_ids, "amount": amount},
        headers=auth,
    )


async def activity(client, auth, name) -> Decimal:
    budget = (await client.get(BUDGET, headers=auth)).json()
    return Decimal(next(ln for ln in budget["lines"] if ln["name"] == name)["activity"])


async def test_a_budget_line_on_a_parent_counts_its_children(client: AsyncClient, auth):
    food = await family(client, auth)
    await spend(client, auth, "Food", "10")
    await spend(client, auth, "Groceries", "20")
    await spend(client, auth, "Dining", "30")
    await spend(client, auth, "Fuel", "40")
    assert (await line(client, auth, "Food", [food["id"]])).status_code == 201
    assert await activity(client, auth, "Food") == 60


async def test_a_line_with_a_parent_and_its_child_counts_the_child_once(client: AsyncClient, auth):
    food = await family(client, auth)
    groceries = next(c for c in await all_categories(client, auth) if c["name"] == "Groceries")
    await spend(client, auth, "Groceries", "20")
    assert (await line(client, auth, "Food", [food["id"], groceries["id"]])).status_code == 201
    assert await activity(client, auth, "Food") == 20


async def test_a_child_is_taken_when_its_parent_is_in_another_line(client: AsyncClient, auth):
    food = await family(client, auth)
    groceries = next(c for c in await all_categories(client, auth) if c["name"] == "Groceries")
    assert (await line(client, auth, "Food", [food["id"]])).status_code == 201
    clash = await line(client, auth, "Weekly shop", [groceries["id"]])
    assert clash.status_code == 409 or clash.status_code == 422
    assert 'already covered by "Food" in the "Food" line' in clash.json()["detail"]


async def test_a_parent_is_taken_when_a_child_is_in_another_line(client: AsyncClient, auth):
    food = await family(client, auth)
    groceries = next(c for c in await all_categories(client, auth) if c["name"] == "Groceries")
    assert (await line(client, auth, "Weekly shop", [groceries["id"]])).status_code == 201
    clash = await line(client, auth, "Food", [food["id"]])
    assert clash.status_code in (409, 422)
    assert 'already covered by "Groceries" in the "Weekly shop" line' in clash.json()["detail"]
    unrelated = next(c for c in await all_categories(client, auth) if c["name"] == "Fuel")
    assert (await line(client, auth, "Car", [unrelated["id"]])).status_code == 201


async def test_a_refund_on_a_child_lowers_the_parents_line(client: AsyncClient, auth):
    food = await family(client, auth)
    bought = await spend(client, auth, "Groceries", "100")
    await line(client, auth, "Food", [food["id"]])
    resp = await client.post(
        TX,
        json={
            "type": "income",
            "amount": "30",
            "transaction_date": DAY,
            "refund_of_id": bought["id"],
        },
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    assert await activity(client, auth, "Food") == 70
