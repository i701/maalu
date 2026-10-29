from datetime import date
from io import BytesIO

from httpx import AsyncClient
from openpyxl import load_workbook

TX = "/api/v1/transactions"
DAY = "2026-03-10"


async def add(client: AsyncClient, auth, **extra) -> dict:
    body = {
        "type": "expense",
        "amount": "10",
        "payment_method": "cash",
        "category": "Food",
        "transaction_date": DAY,
        **extra,
    }
    resp = await client.post(TX, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def test_time_is_optional_and_comes_back(client: AsyncClient, auth):
    plain = await add(client, auth)
    timed = await add(client, auth, transaction_time="14:30:00")
    assert plain["transaction_time"] is None
    assert timed["transaction_time"] == "14:30:00"


async def test_bad_time_is_rejected(client: AsyncClient, auth):
    resp = await client.post(
        TX,
        json={
            "type": "expense",
            "amount": "1",
            "payment_method": "cash",
            "category": "Food",
            "transaction_date": DAY,
            "transaction_time": "25:99",
        },
        headers=auth,
    )
    assert resp.status_code == 422


async def test_update_sets_and_clears_the_time(client: AsyncClient, auth):
    t = await add(client, auth)
    resp = await client.patch(
        f"{TX}/{t['id']}", json={"transaction_time": "09:05:00"}, headers=auth
    )
    assert resp.json()["transaction_time"] == "09:05:00"
    resp = await client.patch(f"{TX}/{t['id']}", json={"transaction_time": None}, headers=auth)
    assert resp.status_code == 200
    assert resp.json()["transaction_time"] is None


async def test_leaving_the_time_out_keeps_it(client: AsyncClient, auth):
    t = await add(client, auth, transaction_time="08:00:00")
    resp = await client.patch(f"{TX}/{t['id']}", json={"description": "x"}, headers=auth)
    assert resp.json()["transaction_time"] == "08:00:00"


async def test_a_day_lists_latest_time_first_and_untimed_last(client: AsyncClient, auth):
    early = await add(client, auth, transaction_time="08:00:00")
    late = await add(client, auth, transaction_time="21:15:00")
    untimed = await add(client, auth)
    later_day = await add(client, auth, transaction_date="2026-03-11")
    resp = await client.get(TX, headers=auth)
    assert [t["id"] for t in resp.json()] == [
        later_day["id"],
        late["id"],
        early["id"],
        untimed["id"],
    ]


async def test_upcoming_reads_earliest_time_first(client: AsyncClient, auth):
    day = date.today().replace(year=date.today().year + 1).isoformat()
    b = await add(client, auth, transaction_date=day, transaction_time="18:00:00", planned=True)
    a = await add(client, auth, transaction_date=day, transaction_time="07:30:00", planned=True)
    resp = await client.get(TX, params={"planned": "true"}, headers=auth)
    assert [t["id"] for t in resp.json()] == [a["id"], b["id"]]


async def test_posting_a_planned_one_can_set_the_time(client: AsyncClient, auth):
    t = await add(client, auth, planned=True)
    resp = await client.post(
        f"{TX}/{t['id']}/post", json={"transaction_time": "12:45:00"}, headers=auth
    )
    assert resp.status_code == 200
    assert resp.json()["transaction_time"] == "12:45:00"
    assert resp.json()["planned"] is False


async def test_export_has_a_time_column(client: AsyncClient, auth):
    await add(client, auth, transaction_time="14:30:00")
    resp = await client.get(f"{TX}/export", headers=auth)
    assert resp.status_code == 200
    rows = list(load_workbook(BytesIO(resp.content))["Transactions"].iter_rows(values_only=True))
    time_col = rows[0].index("Time")
    assert rows[1][time_col].strftime("%H:%M") == "14:30"
