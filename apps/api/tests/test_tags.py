import io
from decimal import Decimal

from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import Signup

TX = "/api/v1/transactions"
TAGS = "/api/v1/tags"
TPL = "/api/v1/templates"


def expense(amount="10", **extra) -> dict:
    return {
        "type": "expense",
        "amount": amount,
        "payment_method": "cash",
        "category": "Misc",
        "transaction_date": "2026-03-01",
        **extra,
    }


async def add(client: AsyncClient, auth, **extra) -> dict:
    resp = await client.post(TX, json=expense(**extra), headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


def names(t: dict) -> list[str]:
    return [tag["name"] for tag in t["tags"]]


async def test_tags_are_found_or_created_by_name(client: AsyncClient, auth):
    t = await add(client, auth, tags=["Trip", "need"])
    assert names(t) == ["need", "Trip"]  # by name, ignoring case
    listed = (await client.get(TAGS, headers=auth)).json()
    assert [x["name"] for x in listed] == ["need", "Trip"]
    # Reused, not duplicated, whatever the case or spacing.
    again = await add(client, auth, tags=["  trip ", "TRIP"])
    assert names(again) == ["Trip"]
    assert len((await client.get(TAGS, headers=auth)).json()) == 2


async def test_no_tags_by_default(client: AsyncClient, auth):
    assert (await add(client, auth))["tags"] == []


async def test_tag_colours_spread_over_the_palette(client: AsyncClient, auth):
    colours = []
    for i in range(9):
        t = await add(client, auth, tags=[f"tag{i}"])
        colours.append(t["tags"][0]["color"])
    assert len(set(colours[:8])) == 8  # every palette colour before any repeats
    assert colours[8] == colours[0]
    # A tag keeps its colour.
    again = await add(client, auth, tags=["tag0"])
    assert again["tags"][0]["color"] == colours[0]


async def test_limits(client: AsyncClient, auth):
    too_many = [f"t{i}" for i in range(11)]
    assert (await client.post(TX, json=expense(tags=too_many), headers=auth)).status_code == 422
    blank = await client.post(TX, json=expense(tags=["  "]), headers=auth)
    assert blank.status_code == 422
    long = await client.post(TX, json=expense(tags=["x" * 51]), headers=auth)
    assert long.status_code == 422


async def test_update_replaces_clears_and_leaves_alone(client: AsyncClient, auth):
    t = await add(client, auth, tags=["a", "b"])
    url = f"{TX}/{t['id']}"
    untouched = await client.patch(url, json={"description": "x"}, headers=auth)
    assert names(untouched.json()) == ["a", "b"]
    replaced = await client.patch(url, json={"tags": ["c"]}, headers=auth)
    assert names(replaced.json()) == ["c"]
    cleared = await client.patch(url, json={"tags": []}, headers=auth)
    assert cleared.json()["tags"] == []
    await client.patch(url, json={"tags": ["a"]}, headers=auth)
    nulled = await client.patch(url, json={"tags": None}, headers=auth)
    assert nulled.json()["tags"] == []
    # The tags themselves stay.
    assert len((await client.get(TAGS, headers=auth)).json()) == 3


async def test_filter_by_tag_everywhere(client: AsyncClient, auth):
    trip = await add(client, auth, amount="40", tags=["Trip"])
    await add(client, auth, amount="5", tags=["Other"])
    await add(client, auth, amount="3")
    tag_id = (await client.get(TAGS, params={"q": "trip"}, headers=auth)).json()[0]["id"]

    listed = await client.get(TX, params={"tag_id": tag_id}, headers=auth)
    assert [t["id"] for t in listed.json()] == [trip["id"]]
    assert listed.headers["X-Total-Count"] == "1"
    summary = (await client.get(f"{TX}/summary", params={"tag_id": tag_id}, headers=auth)).json()
    assert Decimal(summary["expense"]) == 40
    cats = (await client.get(f"{TX}/categories", params={"tag_id": tag_id}, headers=auth)).json()
    assert [Decimal(c["total"]) for c in cats["categories"]] == [Decimal(40)]


async def test_search_matches_tag_names(client: AsyncClient, auth):
    t = await add(client, auth, tags=["Holiday"])
    await add(client, auth)
    found = await client.get(TX, params={"q": "holi"}, headers=auth)
    assert [x["id"] for x in found.json()] == [t["id"]]


async def test_export_has_a_tags_column_and_the_filter_label(client: AsyncClient, auth):
    await add(client, auth, tags=["Trip", "Need"])
    await add(client, auth)
    tag_id = (await client.get(TAGS, params={"q": "trip"}, headers=auth)).json()[0]["id"]
    resp = await client.get(f"{TX}/export", params={"tag_id": tag_id}, headers=auth)
    wb = load_workbook(io.BytesIO(resp.content))
    rows = list(wb["Transactions"].iter_rows(values_only=True))
    header = list(rows[0])
    assert "Tags" in header
    assert len(rows) == 2  # header + the one tagged transaction
    assert rows[1][header.index("Tags")] == "Need, Trip"
    summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]]
    assert any("Tag: Trip" in str(cell) for row in summary for cell in row)


async def test_trashed_transactions_keep_tags_but_do_not_count(client: AsyncClient, auth):
    t = await add(client, auth, tags=["Trip"])
    tag = (await client.get(TAGS, headers=auth)).json()[0]
    assert tag["transaction_count"] == 1
    await client.delete(f"{TX}/{t['id']}", headers=auth)
    assert (await client.get(TAGS, headers=auth)).json()[0]["transaction_count"] == 0
    restored = (await client.post(f"{TX}/{t['id']}/restore", headers=auth)).json()
    assert names(restored) == ["Trip"]
    assert (await client.get(TAGS, headers=auth)).json()[0]["transaction_count"] == 1


async def test_deleting_a_transaction_for_good_keeps_the_tag(client: AsyncClient, auth):
    t = await add(client, auth, tags=["Trip"])
    await client.delete(f"{TX}/{t['id']}", headers=auth)
    await client.delete(f"{TX}/{t['id']}/permanent", headers=auth)
    tags = (await client.get(TAGS, headers=auth)).json()
    assert [x["name"] for x in tags] == ["Trip"]
    assert tags[0]["transaction_count"] == 0


async def test_list_with_counts_and_paging(client: AsyncClient, auth):
    await add(client, auth, tags=["b", "a"])
    await add(client, auth, tags=["a"])
    resp = await client.get(TAGS, headers=auth)
    assert [(t["name"], t["transaction_count"]) for t in resp.json()] == [("a", 2), ("b", 1)]
    assert resp.headers["X-Total-Count"] == "2"
    page = await client.get(TAGS, params={"limit": 1, "offset": 1}, headers=auth)
    assert [t["name"] for t in page.json()] == ["b"]
    assert page.headers["X-Total-Count"] == "2"
    found = await client.get(TAGS, params={"q": "A"}, headers=auth)
    assert [t["name"] for t in found.json()] == ["a"]


async def test_create_rename_and_clashes(client: AsyncClient, auth):
    created = await client.post(TAGS, json={"name": " Trip "}, headers=auth)
    assert created.status_code == 201
    assert created.json()["name"] == "Trip"
    assert (await client.post(TAGS, json={"name": "trip"}, headers=auth)).status_code == 409
    other = (await client.post(TAGS, json={"name": "Need"}, headers=auth)).json()
    renamed = await client.patch(f"{TAGS}/{other['id']}", json={"name": "Wants"}, headers=auth)
    assert renamed.json()["name"] == "Wants"
    clash = await client.patch(f"{TAGS}/{other['id']}", json={"name": "TRIP"}, headers=auth)
    assert clash.status_code == 409
    same = await client.patch(f"{TAGS}/{other['id']}", json={"name": "wants"}, headers=auth)
    assert same.status_code == 200
    # A rename shows on the transactions that carry it.
    t = await add(client, auth, tags=["Wants"])
    await client.patch(f"{TAGS}/{other['id']}", json={"name": "Treats"}, headers=auth)
    assert names((await client.get(f"{TX}/{t['id']}", headers=auth)).json()) == ["Treats"]


async def test_merge_moves_everything_without_duplicates(client: AsyncClient, auth):
    both = await add(client, auth, tags=["Trip", "Holiday"])
    only_source = await add(client, auth, tags=["Holiday"])
    await client.post(TPL, json={"name": "T", "type": "expense", "tags": ["Holiday"]}, headers=auth)
    tags = {t["name"]: t["id"] for t in (await client.get(TAGS, headers=auth)).json()}

    merged = await client.post(
        f"{TAGS}/{tags['Holiday']}/merge", json={"into_id": tags["Trip"]}, headers=auth
    )
    assert merged.status_code == 200
    assert merged.json()["transaction_count"] == 2
    assert merged.json()["template_count"] == 1
    assert [t["name"] for t in (await client.get(TAGS, headers=auth)).json()] == ["Trip"]
    assert names((await client.get(f"{TX}/{both['id']}", headers=auth)).json()) == ["Trip"]
    assert names((await client.get(f"{TX}/{only_source['id']}", headers=auth)).json()) == ["Trip"]
    template = (await client.get(TPL, headers=auth)).json()[0]
    assert [t["name"] for t in template["tags"]] == ["Trip"]


async def test_merge_into_itself_is_refused(client: AsyncClient, auth):
    tag = (await client.post(TAGS, json={"name": "A"}, headers=auth)).json()
    resp = await client.post(f"{TAGS}/{tag['id']}/merge", json={"into_id": tag["id"]}, headers=auth)
    assert resp.status_code == 422


async def test_delete_only_unlinks(client: AsyncClient, auth):
    t = await add(client, auth, tags=["Trip", "Need"])
    trip = next(x for x in t["tags"] if x["name"] == "Trip")
    assert (await client.delete(f"{TAGS}/{trip['id']}", headers=auth)).status_code == 204
    after = (await client.get(f"{TX}/{t['id']}", headers=auth)).json()
    assert names(after) == ["Need"]
    assert (await client.get(f"{TAGS}/{trip['id']}", headers=auth)).status_code == 404


async def test_templates_carry_tags(client: AsyncClient, auth):
    resp = await client.post(
        TPL,
        json={"name": "Trip lunch", "type": "expense", "tags": ["Trip", "trip", "Food"]},
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    assert names(resp.json()) == ["Food", "Trip"]
    assert len((await client.get(TAGS, headers=auth)).json()) == 2
    # Deleting the template leaves the tags.
    await client.delete(f"{TPL}/{resp.json()['id']}", headers=auth)
    assert len((await client.get(TAGS, headers=auth)).json()) == 2


async def test_tags_are_private(client: AsyncClient, auth, signup: Signup):
    t = await add(client, auth, tags=["Trip"])
    tag = t["tags"][0]
    other = await signup("other@example.com")
    assert (await client.get(TAGS, headers=other)).json() == []
    assert (await client.get(f"{TAGS}/{tag['id']}", headers=other)).status_code == 404
    assert (
        await client.patch(f"{TAGS}/{tag['id']}", json={"name": "x"}, headers=other)
    ).status_code == 404
    assert (await client.delete(f"{TAGS}/{tag['id']}", headers=other)).status_code == 404
    mine = (await client.post(TAGS, json={"name": "Mine"}, headers=auth)).json()
    cross = await client.post(
        f"{TAGS}/{tag['id']}/merge", json={"into_id": mine["id"]}, headers=other
    )
    assert cross.status_code == 404
    # Another user can use the same name for their own tag.
    assert (await client.post(TAGS, json={"name": "Trip"}, headers=other)).status_code == 201


async def test_requires_login(client: AsyncClient):
    assert (await client.get(TAGS)).status_code == 401


async def _ids(client: AsyncClient, auth) -> dict[str, str]:
    return {t["name"]: t["id"] for t in (await client.get(TAGS, headers=auth)).json()}


async def test_several_tags_match_any_by_default(client: AsyncClient, auth):
    both = await add(client, auth, amount="1", tags=["Trip", "Need"])
    trip = await add(client, auth, amount="2", tags=["Trip"])
    need = await add(client, auth, amount="4", tags=["Need"])
    await add(client, auth, amount="8", tags=["Other"])
    await add(client, auth, amount="16")
    ids = await _ids(client, auth)
    params = [("tag_id", ids["Trip"]), ("tag_id", ids["Need"])]

    listed = await client.get(TX, params=params, headers=auth)
    assert {t["id"] for t in listed.json()} == {both["id"], trip["id"], need["id"]}
    assert listed.headers["X-Total-Count"] == "3"
    # A transaction with both tags counts once in the totals.
    summary = (await client.get(f"{TX}/summary", params=params, headers=auth)).json()
    assert Decimal(summary["expense"]) == 7
    cats = (await client.get(f"{TX}/categories", params=params, headers=auth)).json()
    assert [Decimal(c["total"]) for c in cats["categories"]] == [Decimal(7)]


async def test_match_all_needs_every_tag(client: AsyncClient, auth):
    both = await add(client, auth, amount="1", tags=["Trip", "Need"])
    await add(client, auth, amount="2", tags=["Trip"])
    await add(client, auth, amount="4", tags=["Need"])
    ids = await _ids(client, auth)
    params = [("tag_id", ids["Trip"]), ("tag_id", ids["Need"]), ("tag_match", "all")]

    listed = await client.get(TX, params=params, headers=auth)
    assert [t["id"] for t in listed.json()] == [both["id"]]
    summary = (await client.get(f"{TX}/summary", params=params, headers=auth)).json()
    assert Decimal(summary["expense"]) == 1


async def test_bad_match_value_is_refused(client: AsyncClient, auth):
    resp = await client.get(TX, params={"tag_match": "some"}, headers=auth)
    assert resp.status_code == 422


async def test_export_names_every_tag_in_the_filter(client: AsyncClient, auth):
    await add(client, auth, tags=["Trip", "Need"])
    ids = await _ids(client, auth)
    for match, joiner in (("any", " or "), ("all", " and ")):
        params = [("tag_id", ids["Need"]), ("tag_id", ids["Trip"]), ("tag_match", match)]
        resp = await client.get(f"{TX}/export", params=params, headers=auth)
        wb = load_workbook(io.BytesIO(resp.content))
        text = [str(c) for row in wb["Summary"].iter_rows(values_only=True) for c in row if c]
        assert any(f"Tag: Need{joiner}Trip" in c for c in text), text
