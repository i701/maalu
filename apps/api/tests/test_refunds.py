from datetime import date
from decimal import Decimal

from app.budget_period import period_bounds
from httpx import AsyncClient

from tests.conftest import Signup

TX = "/api/v1/transactions"
STATS = "/api/v1/stats"
BUDGET = "/api/v1/budget"
CAT = "/api/v1/categories"
TODAY = date.today()
DAY = TODAY.isoformat()

ITEMS = [
    {"name": "Eggs", "quantity": "6", "category": "Groceries", "amount": "60.00"},
    {"name": "Headphones", "category": "Gadgets", "amount": "40.00"},
]


async def buy(client, auth, amount="100", category="Groceries", **extra) -> dict:
    body = {
        "type": "expense",
        "amount": amount,
        "payment_method": "debit_card",
        "category": category,
        "transaction_date": DAY,
        "vendor_name": "Corner Shop",
        **extra,
    }
    resp = await client.post(TX, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def buy_items(client, auth, **extra) -> dict:
    body = {
        "type": "expense",
        "amount": "100",
        "payment_method": "cash",
        "transaction_date": DAY,
        "splits": ITEMS,
        **extra,
    }
    resp = await client.post(TX, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def refund(client, auth, purchase: dict, amount="30", **extra):
    body = {
        "type": "income",
        "amount": amount,
        "transaction_date": DAY,
        "refund_of_id": purchase["id"],
        **extra,
    }
    return await client.post(TX, json=body, headers=auth)


async def refund_ok(client, auth, purchase: dict, amount="30", **extra) -> dict:
    resp = await refund(client, auth, purchase, amount, **extra)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def get(client, auth, t: dict) -> dict:
    return (await client.get(f"{TX}/{t['id']}", headers=auth)).json()


async def summary(client, auth, **params) -> dict:
    return (await client.get(f"{TX}/summary", params=params, headers=auth)).json()


async def by_category(client, auth, **params) -> dict[tuple[str, str], str]:
    body = (await client.get(f"{TX}/categories", params=params, headers=auth)).json()
    return {(c["type"], c["category"]): c["total"] for c in body["categories"]}


# --- making a refund ---------------------------------------------------------------------------


async def test_a_refund_gives_money_back_for_a_purchase(client: AsyncClient, auth):
    purchase = await buy(client, auth)
    body = await refund_ok(client, auth, purchase, "30")
    assert body["type"] == "income"
    assert body["category"] == "Groceries"  # taken from the purchase
    assert body["payment_method"] == "debit_card"  # and so is the method, unless given
    assert body["vendor"]["name"] == "Corner Shop"
    ref = body["refund_of"]
    assert ref["id"] == purchase["id"]
    assert ref["vendor_name"] == "Corner Shop"
    assert ref["amount"] == "100.0000"
    assert ref["category"] == "Groceries"
    assert (await get(client, auth, purchase))["refunded_amount"] == "30.0000"
    assert body["refunded_amount"] == "0.0000"


async def test_a_plain_transaction_is_not_a_refund(client: AsyncClient, auth):
    purchase = await buy(client, auth)
    assert purchase["refund_of"] is None
    assert purchase["refunded_amount"] == "0.0000"


async def test_refunds_add_up_to_the_purchase_and_no_more(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "60")
    too_much = await refund(client, auth, purchase, "50")
    assert too_much.status_code == 422
    assert "Only 40.00 of this purchase is left to refund" in too_much.json()["detail"]
    await refund_ok(client, auth, purchase, "40")  # exactly what is left
    assert (await get(client, auth, purchase))["refunded_amount"] == "100.0000"
    assert (await refund(client, auth, purchase, "0.01")).status_code == 422


async def test_only_your_own_posted_expenses_can_be_refunded(
    client: AsyncClient, auth, signup: Signup
):
    income = (
        await client.post(
            TX,
            json={
                "type": "income",
                "amount": "50",
                "payment_method": "cash",
                "category": "Pay",
                "transaction_date": DAY,
            },
            headers=auth,
        )
    ).json()
    assert (await refund(client, auth, income)).status_code == 422
    planned = await buy(client, auth, planned=True)
    assert (await refund(client, auth, planned)).status_code == 422
    trashed = await buy(client, auth)
    await client.delete(f"{TX}/{trashed['id']}", headers=auth)
    assert (await refund(client, auth, trashed)).status_code == 422
    bob = await signup("bob@example.com")
    theirs = await buy(client, bob)
    assert (await refund(client, auth, theirs)).status_code == 422
    assert (
        await refund(client, auth, {"id": "00000000-0000-0000-0000-000000000000"})
    ).status_code == 422


async def test_a_refund_must_be_income(client: AsyncClient, auth):
    purchase = await buy(client, auth)
    body = {
        "type": "expense",
        "amount": "5",
        "payment_method": "cash",
        "category": "Groceries",
        "transaction_date": DAY,
        "refund_of_id": purchase["id"],
    }
    assert (await client.post(TX, json=body, headers=auth)).status_code == 422


async def test_a_refund_has_the_currency_and_rate_of_the_purchase(client: AsyncClient, auth):
    purchase = await buy(client, auth, "10", currency="USD", exchange_rate="15.4")
    body = await refund_ok(client, auth, purchase, "4")
    assert body["currency"] == "USD"
    assert body["exchange_rate"] == "15.40000000"
    assert body["base_amount"] == "61.6000"  # 4 x 15.4: what the purchase took off is given back
    other = await refund(client, auth, purchase, "1", currency="EUR", exchange_rate="17")
    assert other.status_code == 422
    assert "currency of the purchase" in other.json()["detail"]


async def test_a_refund_can_use_its_own_rate_and_account(client: AsyncClient, auth):
    purchase = await buy(client, auth, "10", currency="USD", exchange_rate="15.4")
    body = await refund_ok(client, auth, purchase, "4", exchange_rate="15.6")
    assert body["base_amount"] == "62.4000"


async def test_a_planned_refund_is_reserved_but_counts_for_nothing_yet(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "70", planned=True)
    assert (await refund(client, auth, purchase, "40")).status_code == 422  # only 30 left
    assert Decimal((await summary(client, auth))["refunded"]) == 0


# --- refunding the items of a purchase -------------------------------------------------------


async def test_a_purchase_with_items_is_refunded_item_by_item(client: AsyncClient, auth):
    purchase = await buy_items(client, auth)
    eggs, phones = purchase["splits"]
    body = await refund_ok(
        client,
        auth,
        purchase,
        "40",
        splits=[
            {
                "name": "ignored",
                "category": "ignored",
                "amount": "40",
                "refund_of_split_id": phones["id"],
            }
        ],
    )
    [line] = body["splits"]
    assert (line["name"], line["category"]) == ("Headphones", "Gadgets")  # from the item
    assert line["refund_of_split_id"] == phones["id"]
    again = (await get(client, auth, purchase))["splits"]
    assert again[0]["refunded_amount"] == "0.0000"
    assert again[1]["refunded_amount"] == "40.0000"
    assert eggs["id"] == again[0]["id"]


async def test_item_refunds_take_money_off_the_items_own_categories(client: AsyncClient, auth):
    purchase = await buy_items(client, auth)
    eggs, phones = purchase["splits"]
    await refund_ok(
        client,
        auth,
        purchase,
        "60",
        splits=[
            {"name": "x", "category": "x", "amount": "25", "refund_of_split_id": eggs["id"]},
            {"name": "x", "category": "x", "amount": "35", "refund_of_split_id": phones["id"]},
        ],
    )
    totals = await by_category(client, auth)
    assert totals[("expense", "Groceries")] == "35.0000"  # 60 less 25
    assert totals[("expense", "Gadgets")] == "5.0000"  # 40 less 35


async def test_item_refund_rules(client: AsyncClient, auth):
    purchase = await buy_items(client, auth)
    eggs, _ = purchase["splits"]

    def line(item, amount):
        return {"name": "x", "category": "x", "amount": amount, "refund_of_split_id": item["id"]}

    needs_items = await refund(client, auth, purchase, "10")
    assert needs_items.status_code == 422
    assert "Choose which items" in needs_items.json()["detail"]
    unnamed = await refund(
        client, auth, purchase, "10", splits=[{"name": "x", "category": "x", "amount": "10"}]
    )
    assert unnamed.status_code == 422
    foreign = await refund(
        client, auth, purchase, "10", splits=[line({"id": purchase["id"]}, "10")]
    )
    assert foreign.status_code == 422
    too_much = await refund(client, auth, purchase, "70", splits=[line(eggs, "70")])
    assert too_much.status_code == 422
    assert 'Only 60.00 of "Eggs" is left to refund' in too_much.json()["detail"]
    twice = await refund(client, auth, purchase, "20", splits=[line(eggs, "10"), line(eggs, "10")])
    assert twice.status_code == 422
    await refund_ok(client, auth, purchase, "50", splits=[line(eggs, "50")])
    left = await refund(client, auth, purchase, "20", splits=[line(eggs, "20")])
    assert left.status_code == 422  # only 10 of the eggs are left
    await refund_ok(client, auth, purchase, "10", splits=[line(eggs, "10")])


async def test_a_purchase_without_items_has_none_to_refund_separately(client: AsyncClient, auth):
    purchase = await buy(client, auth)
    resp = await refund(
        client,
        auth,
        purchase,
        "10",
        splits=[{"name": "x", "category": "x", "amount": "10"}],
    )
    assert resp.status_code == 422


# --- what the reports show ---------------------------------------------------------------------


async def test_a_refund_lowers_spending_and_is_not_income(client: AsyncClient, auth):
    await buy(client, auth, "50", category="Pay-day treat")  # unrelated spending
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    totals = await summary(client, auth)
    assert Decimal(totals["income"]) == 0  # a refund is not income
    assert Decimal(totals["expense"]) == 120  # 150 spent less 30 back
    assert Decimal(totals["refunded"]) == 30


async def test_category_totals_net_the_refund_in_its_own_category(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    await buy(client, auth, "20", category="Fuel")
    await refund_ok(client, auth, purchase, "30")
    totals = await by_category(client, auth)
    assert totals[("expense", "Groceries")] == "70.0000"
    assert totals[("expense", "Fuel")] == "20.0000"
    assert not any(kind == "income" for kind, _ in totals)


async def test_filtering_by_category_nets_refunds_there_too(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    cat = next(c for c in (await client.get(CAT, headers=auth)).json() if c["name"] == "Groceries")
    totals = await summary(client, auth, category_id=cat["id"])
    assert Decimal(totals["expense"]) == 70
    assert Decimal(totals["income"]) == 0


async def test_the_monthly_figures_move_the_refund_from_income_to_spending(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    month = (
        await client.get(f"{STATS}/monthly", params={"months": 1, "to": DAY}, headers=auth)
    ).json()["months"][0]
    assert month["income"] == "0.0000"
    assert month["expense"] == "70.0000"


async def test_a_refund_in_a_later_month_lowers_that_months_spending(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100", transaction_date="2026-01-10")
    await refund_ok(client, auth, purchase, "30", transaction_date="2026-02-05")
    months = (
        await client.get(f"{STATS}/monthly", params={"months": 2, "to": "2026-02-15"}, headers=auth)
    ).json()["months"]
    assert [m["expense"] for m in months] == ["100.0000", "-30.0000"]
    assert [m["income"] for m in months] == ["0.0000", "0.0000"]


async def test_a_refund_lowers_what_the_budget_counts_as_spent(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    cat = next(c for c in (await client.get(CAT, headers=auth)).json() if c["name"] == "Groceries")
    await client.post(
        f"{BUDGET}/lines",
        json={"name": "Food", "category_ids": [cat["id"]], "amount": "200"},
        headers=auth,
    )
    await refund_ok(client, auth, purchase, "30")
    budget = (await client.get(BUDGET, headers=auth)).json()
    start, _ = period_bounds(TODAY, 1)
    assert budget["period"]["start"] == start.isoformat()
    assert (
        Decimal(next(line for line in budget["lines"] if line["name"] == "Food")["activity"]) == 70
    )


async def test_a_refund_still_adds_to_the_account_it_arrives_in(client: AsyncClient, auth):
    account = (
        await client.post(
            "/api/v1/accounts",
            json={"name": "Card", "group": "accounts", "opening_amount": "500"},
            headers=auth,
        )
    ).json()
    purchase = await buy(client, auth, "100", account_id=account["id"])
    await refund_ok(client, auth, purchase, "30", account_id=account["id"])
    accounts = (await client.get("/api/v1/accounts", headers=auth)).json()
    assert next(a for a in accounts if a["id"] == account["id"])["balance"] == "430.0000"


# --- a purchase that has refunds -------------------------------------------------------


async def test_a_purchase_with_refunds_keeps_its_amount_category_and_items(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    url = f"{TX}/{purchase['id']}"
    for change in (
        {"amount": "50"},
        {"category": "Fuel"},
        {"currency": "USD", "exchange_rate": "15"},
        {"planned": True},
        {"type": "income"},
        {"splits": ITEMS},
    ):
        resp = await client.patch(url, json=change, headers=auth)
        assert resp.status_code == 422, change
        assert "has refunds" in resp.json()["detail"]
    ok = await client.patch(
        url, json={"description": "weekly shop", "transaction_date": DAY}, headers=auth
    )
    assert ok.status_code == 200
    assert ok.json()["description"] == "weekly shop"


async def test_sending_the_same_values_back_is_not_a_change(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    resp = await client.patch(
        f"{TX}/{purchase['id']}",
        json={
            "amount": "100",
            "category": "groceries",
            "currency": "MVR",
            "type": "expense",
            "planned": False,
            "description": "edited",
        },
        headers=auth,
    )
    assert resp.status_code == 200, resp.text


async def test_editing_other_things_keeps_the_items_the_refunds_point_at(client: AsyncClient, auth):
    purchase = await buy_items(client, auth)
    eggs, _ = purchase["splits"]
    await refund_ok(
        client,
        auth,
        purchase,
        "10",
        splits=[{"name": "x", "category": "x", "amount": "10", "refund_of_split_id": eggs["id"]}],
    )
    resp = await client.patch(
        f"{TX}/{purchase['id']}", json={"splits": ITEMS, "description": "hi"}, headers=auth
    )
    assert resp.status_code == 200, resp.text
    assert [s["id"] for s in resp.json()["splits"]] == [s["id"] for s in purchase["splits"]]
    assert resp.json()["splits"][0]["refunded_amount"] == "10.0000"
    changed = await client.patch(
        f"{TX}/{purchase['id']}",
        json={"splits": [{**ITEMS[0], "amount": "61"}, ITEMS[1]], "amount": "101"},
        headers=auth,
    )
    assert changed.status_code == 422


async def test_a_purchase_with_refunds_stays_until_they_are_in_the_trash(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    refused = await client.delete(f"{TX}/{purchase['id']}", headers=auth)
    assert refused.status_code == 422
    assert "has refunds" in refused.json()["detail"]
    assert (await client.delete(f"{TX}/{back['id']}", headers=auth)).status_code == 204
    assert (await client.delete(f"{TX}/{purchase['id']}", headers=auth)).status_code == 204
    # a refund in the trash gives nothing back
    assert Decimal((await summary(client, auth))["refunded"]) == 0


async def test_bulk_trash_takes_a_purchase_with_its_refunds_but_not_alone(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    alone = await client.post(f"{TX}/bulk/trash", json={"ids": [purchase["id"]]}, headers=auth)
    assert alone.json()["updated"] == 0
    assert "has refunds" in alone.json()["skipped"][0]["reason"]
    both = await client.post(
        f"{TX}/bulk/trash", json={"ids": [purchase["id"], back["id"]]}, headers=auth
    )
    assert both.json() == {"updated": 2, "skipped": []}


async def test_restoring_a_refund_needs_its_purchase_and_room_left(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    first = await refund_ok(client, auth, purchase, "60")
    await client.delete(f"{TX}/{first['id']}", headers=auth)
    await client.delete(f"{TX}/{purchase['id']}", headers=auth)
    no_purchase = await client.post(f"{TX}/{first['id']}/restore", headers=auth)
    assert no_purchase.status_code == 422
    assert "Restore the purchase" in no_purchase.json()["detail"]
    await client.post(f"{TX}/{purchase['id']}/restore", headers=auth)
    await refund_ok(client, auth, purchase, "60")  # uses the room the trashed one left
    full = await client.post(f"{TX}/{first['id']}/restore", headers=auth)
    assert full.status_code == 422
    assert "left to refund" in full.json()["detail"]


async def test_bulk_restore_brings_back_a_purchase_and_its_refund_together(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    await client.post(f"{TX}/bulk/trash", json={"ids": [purchase["id"], back["id"]]}, headers=auth)
    resp = await client.post(
        f"{TX}/bulk/restore", json={"ids": [back["id"], purchase["id"]]}, headers=auth
    )
    assert resp.json() == {"updated": 2, "skipped": []}
    only = await client.post(f"{TX}/bulk/trash", json={"ids": [back["id"]]}, headers=auth)
    await client.post(f"{TX}/bulk/trash", json={"ids": [purchase["id"]]}, headers=auth)
    resp = await client.post(f"{TX}/bulk/restore", json={"ids": [back["id"]]}, headers=auth)
    assert only.json()["updated"] == 1
    assert resp.json()["updated"] == 0
    assert "Restore the purchase" in resp.json()["skipped"][0]["reason"]


async def test_deleting_a_purchase_for_good_turns_its_refunds_into_plain_income(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    await client.delete(f"{TX}/{back['id']}", headers=auth)
    await client.delete(f"{TX}/{purchase['id']}", headers=auth)
    await client.delete(f"{TX}/{purchase['id']}/permanent", headers=auth)
    restored = await client.post(f"{TX}/{back['id']}/restore", headers=auth)
    assert restored.status_code == 200
    assert restored.json()["refund_of"] is None


# --- changing a refund -----------------------------------------------------------------


async def test_a_refunds_amount_can_change_within_what_is_left(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    other = await refund_ok(client, auth, purchase, "50")
    url = f"{TX}/{back['id']}"
    assert (await client.patch(url, json={"amount": "50"}, headers=auth)).status_code == 200
    over = await client.patch(url, json={"amount": "51"}, headers=auth)  # 50 is the other's
    assert over.status_code == 422
    assert "Only 50.00" in over.json()["detail"]
    assert (await get(client, auth, purchase))["refunded_amount"] == "100.0000"
    assert other["id"] != back["id"]


async def test_a_refunds_category_and_type_follow_the_purchase(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    url = f"{TX}/{back['id']}"
    assert (await client.patch(url, json={"category": "Fuel"}, headers=auth)).status_code == 422
    same = await client.patch(url, json={"category": "groceries", "description": "x"}, headers=auth)
    assert same.status_code == 200
    assert (await client.patch(url, json={"type": "expense"}, headers=auth)).status_code == 422


async def test_a_refund_cannot_move_to_another_purchase(client: AsyncClient, auth):
    one = await buy(client, auth, "100")
    two = await buy(client, auth, "100")
    back = await refund_ok(client, auth, one, "30")
    moved = await client.patch(f"{TX}/{back['id']}", json={"refund_of_id": two["id"]}, headers=auth)
    assert moved.status_code == 422


async def test_unlinking_turns_a_refund_into_plain_income(client: AsyncClient, auth):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    resp = await client.patch(f"{TX}/{back['id']}", json={"refund_of_id": None}, headers=auth)
    assert resp.status_code == 200, resp.text
    assert resp.json()["refund_of"] is None
    totals = await summary(client, auth)
    assert Decimal(totals["income"]) == 30  # now it is income
    assert Decimal(totals["expense"]) == 100
    assert (await get(client, auth, purchase))["refunded_amount"] == "0.0000"


async def test_changing_a_refunds_items(client: AsyncClient, auth):
    purchase = await buy_items(client, auth)
    eggs, phones = purchase["splits"]
    back = await refund_ok(
        client,
        auth,
        purchase,
        "10",
        splits=[{"name": "x", "category": "x", "amount": "10", "refund_of_split_id": eggs["id"]}],
    )
    resp = await client.patch(
        f"{TX}/{back['id']}",
        json={
            "amount": "15",
            "splits": [
                {"name": "x", "category": "x", "amount": "15", "refund_of_split_id": phones["id"]}
            ],
        },
        headers=auth,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["splits"][0]["name"] == "Headphones"
    assert (await by_category(client, auth))[("expense", "Gadgets")] == "25.0000"


# --- the rest ---------------------------------------------------------------------------


async def test_bulk_category_changes_leave_refunds_and_refunded_purchases_alone(
    client: AsyncClient, auth
):
    purchase = await buy(client, auth, "100")
    back = await refund_ok(client, auth, purchase, "30")
    plain = await buy(client, auth, "10")
    resp = await client.patch(
        f"{TX}/bulk",
        json={"ids": [purchase["id"], back["id"], plain["id"]], "category": "Fuel"},
        headers=auth,
    )
    assert resp.json()["updated"] == 1
    reasons = {s["id"]: s["reason"] for s in resp.json()["skipped"]}
    assert "keeps its category" in reasons[purchase["id"]]
    assert "from the purchase" in reasons[back["id"]]


async def test_the_export_calls_a_refund_a_refund_and_nets_it(client: AsyncClient, auth):
    import io

    from openpyxl import load_workbook

    purchase = await buy(client, auth, "100")
    await refund_ok(client, auth, purchase, "30")
    resp = await client.get(f"{TX}/export", headers=auth)
    wb = load_workbook(io.BytesIO(resp.content))
    rows = list(wb["Transactions"].iter_rows(values_only=True))
    assert sorted(r[2] for r in rows[1:]) == ["Expense", "Refund"]
    summary_sheet = dict(wb["Summary"].iter_rows(values_only=True))
    assert summary_sheet["Expense (MVR)"] == 70
    assert summary_sheet["Income (MVR)"] == 0
