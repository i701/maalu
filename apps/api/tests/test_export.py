import io
from datetime import date

import pytest
from app import export as export_module
from app.export import XLSX_MEDIA_TYPE
from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import Signup

TX = "/api/v1/transactions"
EXPORT = "/api/v1/transactions/export"
ACC = "/api/v1/accounts"


def body(kind="expense", amount="10", **extra) -> dict:
    return {
        "type": kind,
        "amount": amount,
        "payment_method": "debit_card",
        "category": "Misc",
        "transaction_date": "2026-09-01",
        **extra,
    }


async def download(client, auth, **params):
    resp = await client.get(EXPORT, params=params, headers=auth)
    assert resp.status_code == 200, resp.text
    return resp, load_workbook(io.BytesIO(resp.content))


def rows(wb) -> list[tuple]:
    return list(wb["Transactions"].iter_rows(values_only=True))


async def test_downloads_an_xlsx_with_sensible_headers(client: AsyncClient, auth):
    await client.post(TX, json=body(), headers=auth)
    resp, wb = await download(client, auth)
    assert resp.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert f"maalu-transactions-{date.today().isoformat()}.xlsx" in disposition
    assert wb.sheetnames == ["Transactions", "By category", "Summary"]
    header = rows(wb)[0]
    assert header[0] == "Date"
    assert "Amount (MVR)" in header
    assert "Signed (MVR)" in header


async def test_cells_are_real_numbers_and_dates_not_text(client: AsyncClient, auth):
    await client.post(
        TX,
        json=body(amount="1234.5", category="Rent", vendor_name="Landlord", description="Sept"),
        headers=auth,
    )
    _, wb = await download(client, auth)
    header, row = rows(wb)
    cells = dict(zip(header, row, strict=True))
    assert cells["Category"] == "Rent"
    assert cells["Vendor"] == "Landlord"
    assert cells["Type"] == "Expense"
    assert cells["Payment method"] == "Debit card"
    assert isinstance(cells["Amount"], float)
    assert cells["Amount"] == pytest.approx(1234.5)
    assert cells["Signed (MVR)"] == pytest.approx(-1234.5)  # expenses are negative
    assert cells["Date"].date() == date(2026, 9, 1)  # a real date cell


async def test_income_signs_and_foreign_currency_use_the_base_amount(client: AsyncClient, auth):
    await client.post(TX, json=body("income", "100"), headers=auth)
    foreign = body("expense", "10", currency="USD", exchange_rate="15.4")
    await client.post(TX, json=foreign, headers=auth)
    _, wb = await download(client, auth)
    header, *data = rows(wb)
    by_type = {r[header.index("Type")]: dict(zip(header, r, strict=True)) for r in data}
    assert by_type["Income"]["Signed (MVR)"] == pytest.approx(100)
    usd = by_type["Expense"]
    assert usd["Currency"] == "USD"
    assert usd["Amount"] == pytest.approx(10)
    assert usd["Rate"] == pytest.approx(15.4)
    assert usd["Amount (MVR)"] == pytest.approx(154)
    assert usd["Signed (MVR)"] == pytest.approx(-154)


async def test_transfers_show_both_accounts_and_do_not_count_as_spending(client: AsyncClient, auth):
    a = (
        await client.post(
            ACC, json={"group": "cash", "name": "Wallet", "opening_amount": "500"}, headers=auth
        )
    ).json()
    b = (await client.post(ACC, json={"group": "savings", "name": "Savings"}, headers=auth)).json()
    transfer = {
        "type": "transfer",
        "amount": "200",
        "account_id": a["id"],
        "to_account_id": b["id"],
        "transaction_date": "2026-09-02",
    }
    await client.post(TX, json=transfer, headers=auth)
    _, wb = await download(client, auth)
    header, row = rows(wb)
    cells = dict(zip(header, row, strict=True))
    assert cells["Type"] == "Transfer"
    assert cells["Account"] == "Wallet"
    assert cells["To account"] == "Savings"
    assert cells["Received"] == pytest.approx(200)
    assert cells["Signed (MVR)"] is None  # money only moved
    assert cells["Category"] is None
    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert summary["Expense (MVR)"] == 0
    assert summary["Income (MVR)"] == 0


async def test_exports_follow_the_same_filters_as_the_list(client: AsyncClient, auth):
    await client.post(TX, json=body(category="Lunch", amount="40"), headers=auth)
    await client.post(TX, json=body(category="Rent", amount="500"), headers=auth)
    await client.post(TX, json=body("income", "900", category="Pay"), headers=auth)

    _, only_expenses = await download(client, auth, type="expense")
    assert len(rows(only_expenses)) == 3  # header + 2
    _, searched = await download(client, auth, q="lunch")
    assert [r[2] for r in rows(searched)[1:]] == ["Lunch"]
    _, ranged = await download(client, auth, min_amount="100", max_amount="600")
    assert [r[2] for r in rows(ranged)[1:]] == ["Rent"]
    _, none = await download(client, auth, q="zzz")
    assert len(rows(none)) == 1  # just the header: still a valid file


async def test_summary_sheet_records_filters_and_totals(client: AsyncClient, auth):
    await client.post(TX, json=body(category="Lunch", amount="40"), headers=auth)
    await client.post(TX, json=body(category="Lunch", amount="60"), headers=auth)
    await client.post(TX, json=body("income", "1000", category="Lunch money"), headers=auth)
    _, wb = await download(client, auth, q="lunch", type="expense")
    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert "Search: lunch" in summary["Filters"]
    assert "Type: expense" in summary["Filters"]
    assert summary["Transactions"] == 2
    assert summary["Expense (MVR)"] == pytest.approx(100)
    assert summary["Income (MVR)"] == 0
    assert summary["Base currency"] == "MVR"

    _, unfiltered = await download(client, auth)
    assert "None" in dict(unfiltered["Summary"].iter_rows(values_only=True))["Filters"]


async def test_oldest_first(client: AsyncClient, auth):
    await client.post(TX, json=body(category="Later", transaction_date="2026-09-20"), headers=auth)
    await client.post(
        TX, json=body(category="Earlier", transaction_date="2026-09-01"), headers=auth
    )
    _, wb = await download(client, auth)
    assert [r[2] for r in rows(wb)[1:]] == ["Earlier", "Later"]


async def test_refuses_oversized_exports_instead_of_truncating(
    client: AsyncClient, auth, monkeypatch
):
    monkeypatch.setattr("app.services.transaction.EXPORT_LIMIT", 2)
    for n in range(3):
        await client.post(TX, json=body(category=f"C{n}"), headers=auth)
    resp = await client.get(EXPORT, headers=auth)
    assert resp.status_code == 422
    assert "Narrow the filters" in resp.json()["detail"]
    narrowed = await client.get(EXPORT, params={"q": "C0"}, headers=auth)
    assert narrowed.status_code == 200


async def test_validates_filters_and_requires_auth(client: AsyncClient, auth):
    assert (await client.get(EXPORT)).status_code == 401
    backwards = await client.get(
        EXPORT, params={"min_amount": "9", "max_amount": "1"}, headers=auth
    )
    assert backwards.status_code == 422
    assert (await client.get(EXPORT, params={"type": "refund"}, headers=auth)).status_code == 422


async def test_only_exports_your_own_transactions(client: AsyncClient, auth, signup: Signup):
    await client.post(TX, json=body(category="Mine"), headers=auth)
    bob = await signup("bob@example.com")
    await client.post(TX, json=body(category="Bobs"), headers=bob)
    _, wb = await download(client, auth)
    assert [r[2] for r in rows(wb)[1:]] == ["Mine"]


def test_workbook_builder_handles_no_rows():
    data = export_module.build_workbook(
        [], base_currency="MVR", filter_labels=[], generated_on=date(2026, 10, 10)
    )
    wb = load_workbook(io.BytesIO(data))
    assert len(list(wb["Transactions"].iter_rows(values_only=True))) == 1


async def test_split_transactions_get_an_items_sheet_and_category_report(client: AsyncClient, auth):
    lines = [
        {"name": "Eggs", "quantity": "6", "category": "Groceries", "amount": "60"},
        {"name": "Snickers", "category": "Junk food", "amount": "10.50", "tax_amount": "0.78"},
    ]
    await client.post(TX, json=body(amount="70.50", splits=lines), headers=auth)
    await client.post(TX, json=body(amount="30", category="Groceries"), headers=auth)
    _, wb = await download(client, auth)
    assert wb.sheetnames == ["Transactions", "Items", "By category", "Summary"]

    header, *data = rows(wb)
    split_row = next(r for r in data if r[header.index("Amount")] == pytest.approx(70.5))
    assert split_row[header.index("Category")] == "Groceries, Junk food"

    items_header, *items = list(wb["Items"].iter_rows(values_only=True))
    by_item = {
        r[items_header.index("Item")]: dict(zip(items_header, r, strict=True)) for r in items
    }
    assert by_item["Eggs"]["Quantity"] == pytest.approx(6)
    assert by_item["Snickers"]["GST"] == pytest.approx(0.78)
    assert by_item["Snickers"]["Category"] == "Junk food"

    report = {(r[0], r[1]): r for r in list(wb["By category"].iter_rows(values_only=True))[1:]}
    assert report[("Expense", "Groceries")][3] == pytest.approx(90)  # 60 + the plain 30
    assert report[("Expense", "Groceries")][2] == 2
    assert report[("Expense", "Junk food")][3] == pytest.approx(10.5)
