import io
from datetime import date, timedelta
from decimal import Decimal

from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import Signup

TX = "/api/v1/transactions"
ACC = "/api/v1/accounts"
BUDGET = "/api/v1/budget"
STATS = "/api/v1/stats"
CAT = "/api/v1/categories"
GOALS = "/api/v1/goals"
TODAY = date.today()


async def add(client: AsyncClient, auth, **extra) -> dict:
    body = {
        "type": "expense",
        "amount": "10",
        "payment_method": "cash",
        "category": "Rent",
        "transaction_date": (TODAY + timedelta(days=5)).isoformat(),
        **extra,
    }
    resp = await client.post(TX, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def planned(client: AsyncClient, auth, **extra) -> dict:
    return await add(client, auth, planned=True, **extra)


async def account(client: AsyncClient, auth, name="Main", **extra) -> dict:
    body = {"group": "accounts", "name": name, "opening_amount": "100", **extra}
    resp = await client.post(ACC, json=body, headers=auth)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def upcoming(client: AsyncClient, auth) -> list[dict]:
    return (await client.get(TX, params={"planned": True}, headers=auth)).json()


async def test_a_planned_transaction_is_flagged_and_defaults_to_posted(client: AsyncClient, auth):
    p = await planned(client, auth)
    assert p["planned"] is True
    assert (await add(client, auth))["planned"] is False


async def test_planned_ones_have_their_own_list(client: AsyncClient, auth):
    p = await planned(client, auth)
    posted = await add(client, auth, transaction_date=TODAY.isoformat())
    main = await client.get(TX, headers=auth)
    assert [t["id"] for t in main.json()] == [posted["id"]]
    assert main.headers["X-Total-Count"] == "1"
    soon = await client.get(TX, params={"planned": True}, headers=auth)
    assert [t["id"] for t in soon.json()] == [p["id"]]
    assert soon.headers["X-Total-Count"] == "1"


async def test_upcoming_reads_soonest_first_and_includes_overdue(client: AsyncClient, auth):
    later = await planned(client, auth, transaction_date=(TODAY + timedelta(days=20)).isoformat())
    overdue = await planned(client, auth, transaction_date=(TODAY - timedelta(days=3)).isoformat())
    soon = await planned(client, auth, transaction_date=(TODAY + timedelta(days=2)).isoformat())
    ids = [t["id"] for t in await upcoming(client, auth)]
    assert ids == [overdue["id"], soon["id"], later["id"]]


async def test_planned_counts_toward_no_totals(client: AsyncClient, auth):
    await planned(client, auth, amount="500")
    await add(client, auth, amount="20", transaction_date=TODAY.isoformat())
    summary = (await client.get(f"{TX}/summary", headers=auth)).json()
    assert Decimal(summary["expense"]) == 20
    cats = (await client.get(f"{TX}/categories", headers=auth)).json()["categories"]
    assert [Decimal(c["total"]) for c in cats] == [Decimal(20)]
    exported = await client.get(f"{TX}/export", headers=auth)
    rows = list(
        load_workbook(io.BytesIO(exported.content))["Transactions"].iter_rows(values_only=True)
    )
    assert len(rows) == 2  # header and the one posted transaction


async def test_an_account_balance_ignores_planned_until_posted(client: AsyncClient, auth):
    a = await account(client, auth)
    p = await planned(client, auth, amount="30", account_id=a["id"])
    row = (await client.get(f"{ACC}/{a['id']}", headers=auth)).json()
    assert Decimal(row["balance"]) == 100
    assert row["transaction_count"] == 0
    summary = (await client.get(f"{ACC}/summary", headers=auth)).json()
    assert Decimal(summary["assets"]) == 100

    posted = await client.post(f"{TX}/{p['id']}/post", headers=auth)
    assert posted.status_code == 200, posted.text
    row = (await client.get(f"{ACC}/{a['id']}", headers=auth)).json()
    assert Decimal(row["balance"]) == 70
    assert row["transaction_count"] == 1


async def test_the_balance_trend_and_monthly_stats_ignore_planned(client: AsyncClient, auth):
    await account(client, auth)
    await planned(client, auth, amount="40", transaction_date=TODAY.isoformat())
    monthly = (await client.get(f"{STATS}/monthly", params={"months": 1}, headers=auth)).json()
    assert all(Decimal(m["expense"]) == 0 for m in monthly["months"])
    day = TODAY.isoformat()
    trend = (
        await client.get(f"{STATS}/balance-trend", params={"from": day, "to": day}, headers=auth)
    ).json()
    assert Decimal(trend["points"][0]["total"]) == 100


async def test_a_budget_ignores_planned_until_posted(client: AsyncClient, auth):
    p = await planned(client, auth, amount="60", transaction_date=TODAY.isoformat())
    cat = next(c["id"] for c in (await client.get(CAT, headers=auth)).json() if c["name"] == "Rent")
    await client.post(
        f"{BUDGET}/lines",
        json={"name": "Rent", "category_ids": [cat], "amount": "100"},
        headers=auth,
    )

    def spent(budget: dict) -> Decimal:
        return Decimal(budget["lines"][0]["activity"])

    assert spent((await client.get(BUDGET, headers=auth)).json()) == 0
    await client.post(f"{TX}/{p['id']}/post", headers=auth)
    assert spent((await client.get(BUDGET, headers=auth)).json()) == 60


async def test_post_keeps_the_plan_unless_told_otherwise(client: AsyncClient, auth):
    day = TODAY + timedelta(days=5)
    p = await planned(client, auth, amount="100", transaction_date=day.isoformat())
    posted = (await client.post(f"{TX}/{p['id']}/post", headers=auth)).json()
    assert posted["planned"] is False
    assert posted["transaction_date"] == day.isoformat()
    assert Decimal(posted["amount"]) == 100


async def test_post_can_correct_the_date_and_amount(client: AsyncClient, auth):
    p = await planned(client, auth, amount="100")
    posted = await client.post(
        f"{TX}/{p['id']}/post",
        json={"transaction_date": TODAY.isoformat(), "amount": "92.5"},
        headers=auth,
    )
    body = posted.json()
    assert body["transaction_date"] == TODAY.isoformat()
    assert Decimal(body["amount"]) == Decimal("92.5")
    assert Decimal(body["base_amount"]) == Decimal("92.5")
    assert body["planned"] is False


async def test_post_recomputes_the_converted_amount_for_a_foreign_currency(
    client: AsyncClient, auth
):
    p = await planned(client, auth, amount="10", currency="USD", exchange_rate="15")
    body = (await client.post(f"{TX}/{p['id']}/post", json={"amount": "12"}, headers=auth)).json()
    assert Decimal(body["base_amount"]) == 180  # 12 x 15


async def test_only_a_planned_one_can_be_posted(client: AsyncClient, auth):
    t = await add(client, auth, transaction_date=TODAY.isoformat())
    assert (await client.post(f"{TX}/{t['id']}/post", headers=auth)).status_code == 409
    missing = await client.post(f"{TX}/00000000-0000-0000-0000-000000000000/post", headers=auth)
    assert missing.status_code == 404


async def test_post_is_private_and_skips_the_trash(client: AsyncClient, auth, signup: Signup):
    p = await planned(client, auth)
    other = await signup("other@example.com")
    assert (await client.post(f"{TX}/{p['id']}/post", headers=other)).status_code == 404
    await client.delete(f"{TX}/{p['id']}", headers=auth)
    assert (await client.post(f"{TX}/{p['id']}/post", headers=auth)).status_code == 404


async def test_the_flag_can_be_changed_by_an_edit(client: AsyncClient, auth):
    t = await add(client, auth, transaction_date=TODAY.isoformat())
    url = f"{TX}/{t['id']}"
    assert (await client.patch(url, json={"planned": True}, headers=auth)).json()["planned"] is True
    assert [x["id"] for x in await upcoming(client, auth)] == [t["id"]]
    back = await client.patch(url, json={"planned": False}, headers=auth)
    assert back.json()["planned"] is False
    assert await upcoming(client, auth) == []
    assert (await client.patch(url, json={"planned": None}, headers=auth)).status_code == 422


async def test_trashing_and_restoring_keeps_it_planned(client: AsyncClient, auth):
    p = await planned(client, auth)
    await client.delete(f"{TX}/{p['id']}", headers=auth)
    assert await upcoming(client, auth) == []
    trash = (await client.get(f"{TX}/trash", headers=auth)).json()
    assert [t["id"] for t in trash] == [p["id"]]
    restored = (await client.post(f"{TX}/{p['id']}/restore", headers=auth)).json()
    assert restored["planned"] is True
    assert [t["id"] for t in await upcoming(client, auth)] == [p["id"]]


async def test_bulk_actions_reach_planned_ones(client: AsyncClient, auth):
    p = await planned(client, auth)
    edit = await client.patch(
        f"{TX}/bulk", json={"ids": [p["id"]], "add_tags": ["Due"]}, headers=auth
    )
    assert edit.json()["updated"] == 1
    trash = await client.post(f"{TX}/bulk/trash", json={"ids": [p["id"]]}, headers=auth)
    assert trash.json()["updated"] == 1
    assert await upcoming(client, auth) == []


async def test_filters_still_apply_to_the_upcoming_list(client: AsyncClient, auth):
    a = await planned(client, auth, category="Rent", amount="800")
    await planned(client, auth, category="Insurance", amount="50")
    big = await client.get(TX, params={"planned": True, "min_amount": "100"}, headers=auth)
    assert [t["id"] for t in big.json()] == [a["id"]]


async def test_a_goal_ignores_planned_transfers(client: AsyncClient, auth):
    src = await account(client, auth, "Main", opening_amount="1000")
    dst = await account(client, auth, "Savings", group="savings", opening_amount="0")
    goal = await client.post(
        GOALS,
        json={
            "name": "Trip",
            "target_amount": "500",
            "account_id": dst["id"],
            "monthly_amount": "100",
        },
        headers=auth,
    )
    assert goal.status_code == 201, goal.text
    await client.post(
        TX,
        json={
            "type": "transfer",
            "amount": "100",
            "transaction_date": TODAY.isoformat(),
            "account_id": src["id"],
            "to_account_id": dst["id"],
            "planned": True,
        },
        headers=auth,
    )
    row = (await client.get(GOALS, headers=auth)).json()[0]
    assert Decimal(row["saved"]) == 0
