import io
from decimal import Decimal

from httpx import AsyncClient
from openpyxl import load_workbook

TX = "/api/v1/transactions"
SUMMARY = "/api/v1/transactions/summary"
EXPORT = "/api/v1/transactions/export"
TOTAL = "x-total-count"


def tx(day: str, kind="expense", amount="10", **extra) -> dict:
    return {
        "type": kind,
        "amount": amount,
        "payment_method": "cash",
        "category": f"On {day}",
        "transaction_date": day,
        **extra,
    }


async def days(client, auth, **params) -> list[str]:
    resp = await client.get(TX, params=params, headers=auth)
    assert resp.status_code == 200, resp.text
    return sorted(t["transaction_date"] for t in resp.json())


async def seed(client, auth):
    for day in ("2026-08-31", "2026-09-01", "2026-09-15", "2026-09-30", "2026-10-01"):
        await client.post(TX, json=tx(day), headers=auth)


async def test_range_is_inclusive_at_both_ends(client: AsyncClient, auth):
    await seed(client, auth)
    september = await days(client, auth, date_from="2026-09-01", date_to="2026-09-30")
    assert september == ["2026-09-01", "2026-09-15", "2026-09-30"]
    assert await days(client, auth, date_from="2026-09-15", date_to="2026-09-15") == ["2026-09-15"]


async def test_open_ended_ranges(client: AsyncClient, auth):
    await seed(client, auth)
    assert await days(client, auth, date_from="2026-09-30") == ["2026-09-30", "2026-10-01"]
    assert await days(client, auth, date_to="2026-08-31") == ["2026-08-31"]
    assert await days(client, auth, date_from="2027-01-01") == []


async def test_range_drives_the_pager_total_and_combines_with_other_filters(
    client: AsyncClient, auth
):
    await seed(client, auth)
    await client.post(TX, json=tx("2026-09-10", "income", "500", category="Pay"), headers=auth)

    page = await client.get(
        TX, params={"date_from": "2026-09-01", "date_to": "2026-09-30", "limit": 2}, headers=auth
    )
    assert len(page.json()) == 2
    assert page.headers[TOTAL] == "4"  # three expenses + the income, ignoring the page size

    expenses = await client.get(
        TX,
        params={"date_from": "2026-09-01", "date_to": "2026-09-30", "type": "expense"},
        headers=auth,
    )
    assert expenses.headers[TOTAL] == "3"
    searched = await client.get(TX, params={"date_from": "2026-09-01", "q": "pay"}, headers=auth)
    assert [t["category"] for t in searched.json()] == ["Pay"]


async def test_totals_follow_the_date_range(client: AsyncClient, auth):
    await client.post(TX, json=tx("2026-08-10", amount="100"), headers=auth)
    await client.post(TX, json=tx("2026-09-10", amount="40"), headers=auth)
    await client.post(TX, json=tx("2026-09-20", "income", "700"), headers=auth)
    data = (
        await client.get(
            SUMMARY, params={"date_from": "2026-09-01", "date_to": "2026-09-30"}, headers=auth
        )
    ).json()
    assert Decimal(data["expense"]) == 40  # August's 100 is outside the range
    assert Decimal(data["income"]) == 700


async def test_export_follows_the_date_range_and_records_it(client: AsyncClient, auth):
    await seed(client, auth)
    resp = await client.get(
        EXPORT, params={"date_from": "2026-09-01", "date_to": "2026-09-30"}, headers=auth
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    dates = [
        r[0].date().isoformat() for r in list(wb["Transactions"].iter_rows(values_only=True))[1:]
    ]
    assert dates == ["2026-09-01", "2026-09-15", "2026-09-30"]
    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert "Dates: 2026-09-01 to 2026-09-30" in summary["Filters"]


async def test_invalid_dates_are_rejected(client: AsyncClient, auth):
    backwards = await client.get(
        TX, params={"date_from": "2026-10-01", "date_to": "2026-09-01"}, headers=auth
    )
    assert backwards.status_code == 422
    assert "date_from" in backwards.json()["detail"]
    for bad in ({"date_from": "yesterday"}, {"date_to": "2026-13-45"}):
        assert (await client.get(TX, params=bad, headers=auth)).status_code == 422, bad
    for path in (SUMMARY, EXPORT):
        resp = await client.get(
            path, params={"date_from": "2026-10-01", "date_to": "2026-09-01"}, headers=auth
        )
        assert resp.status_code == 422, path
